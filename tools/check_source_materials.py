#!/usr/bin/env python3
"""Optional source audit. Requires the private Wytyczne directory and PyMuPDF.
Never writes the inputs. Run before accepting editorial/content changes.
"""
import argparse
import json
from collections import Counter
from hashlib import sha256
from pathlib import Path
import re

import pymupdf
from openpyxl import load_workbook
from make_content_template import Document, ROOT, normal


def tokens(text):
    return ' '.join(re.findall(r'[^\W_]+|[<>…]', text))


def corrected(text):
    for before, after in {'MTQPlus': 'MTQ Plus', 'Reflecitve listening': 'Reflective listening', 'Prioritatization': 'Prioritization', 'Interwention': 'Intervention', 'Dysposition': 'Disposition'}.items():
        text = text.replace(before, after)
    return re.sub(r'\bMTQ\b(?!\s+Plus)', 'MTQ Plus', text)


def audit(source):
    hashes = json.loads((ROOT / 'docs/source-hashes.json').read_text())
    for filename, expected in hashes.items():
        assert sha256((source / filename).read_bytes()).hexdigest() == expected, ('Zmienione źródło', filename)
    print(f'OK — {len(hashes)} oryginałów, zgodne SHA-256; źródła pozostają niezmienione.')
    wb = load_workbook(source / 'teksty-strony (2) (1).xlsx', read_only=True, data_only=True)
    rows = [r for r in wb['Teksty'].iter_rows(min_row=2, values_only=True) if r[0]]
    wb.close()
    assert len(rows) == 109
    assert Counter(bool(r[5]) for r in rows)[True] == 45
    assert sum(r[5] == 'USUŃ' for r in rows) == 16
    assert sum(bool(r[6]) for r in rows) == 16
    home = Document((ROOT / 'index.html').read_text())
    for r in rows:
        ident, before, after = r[0], r[3], r[5]
        nodes = [n for n in home.nodes if n.attrs.get('data-source-id') == ident]
        if after == 'USUŃ':
            assert not nodes, ('Niewykonane usunięcie', ident)
            continue
        expected = normal(corrected(after or before))
        if ident in {'T089', 'T091', 'T093', 'T106'}:
            target = {'T089': 'T090', 'T091': 'T092', 'T093': 'T094', 'T106': 'T107'}[ident]
            node = next(n for n in home.nodes if n.attrs.get('data-source-id') == target)
            assert node.attrs['href'] == expected
        elif ident == 'T003':
            assert any(n.tag == 'a' and n.text() == expected and n.attrs.get('href') == '#main-content' for n in home.nodes)
        elif ident == 'T088':
            assert nodes[0].text() == expected.split(' CENNIK ')[0]
            assert 'Jeśli jedyną przeszkodą są pieniądze, zadzwoń, napisz, razem znajdziemy najlepsze dla Ciebie rozwiązanie.' in home.root.text()
            # Semantic price terms/values are independently checked by check_site.py.
        else:
            assert len(nodes) == 1, (ident, len(nodes))
            node = nodes[0]
            actual = node.attrs.get('content') if node.tag == 'meta' else node.attrs.get('alt') if node.tag == 'img' else node.text()
            assert actual == expected, (ident, actual, expected)
    print('OK — T001–T109: 45 wartości F, 16 usunięć, zachowane puste F; kontrakt cennika sprawdzany osobno.')
    routes = {32: 'moja-droga.html', 36: 'twoja-droga.html', 40: 'warsztaty-i-szkolenia.html', 43: 'interwencja-kryzysowa.html', 46: 'coaching.html', 49: 'terapia-dzwiekiem.html', 52: 'mtq-plus.html', 55: 'prism-brain-mapping.html'}
    for ident, route in routes.items():
        assert next(n for n in home.nodes if n.attrs.get('data-source-id') == f'T{ident:03}').attrs['href'] == route
    assert home.root.text().find('Zmiana zaczyna się od samoświadomości.') >= 0
    print('OK — instrukcje G: osiem tras; kolejność i sekcje sprawdzane przez check_site.py.')
    pdf = source / 'Brief stronka.pdf'
    assert sha256(pdf.read_bytes()).hexdigest() == 'bb41337729b3c86bb3c3b229c5d82ac04a49d2efac32cb9a8c3a08849b0e47f8', 'Brief zmienił się; zaktualizuj mapowanie po przeglądzie źródła.'
    blocks = []
    with pymupdf.open(pdf) as document:
        assert len(document) == 19
        for page in document:
            for block in page.get_text('dict')['blocks']:
                if 'lines' in block:
                    text = ''.join(s['text'] for line in block['lines'] for s in line['spans'])
                    blocks.append(normal(text.replace('\u200b', '')))
    ranges = {
        'index.html': list(range(51, 57)),
        'moja-droga.html': [67, 69, 70, 71, 72, 75, 76, 78, 79, 81, 82, 83, 85, 87, 88],
        'twoja-droga.html': [92, 93, 94, 96, 97, 98, 99, 101, 102, 103, 105],
        'coaching.html': [108, 110, 112, 113, 114, 116, 117, 118, 120, 121, 122, 123, 125, 126, 127, 128],
        'interwencja-kryzysowa.html': [132, 133, 134, 135, 136, 138, 139, 140, 141, 143, 144, 145, 146, 147, 149, 150, 152, 153, 154, 155, 156, 157, 158, 159, 160],
        'prism-brain-mapping.html': [164, 166, 168, 169, 170, 173, 174, 175, 176, 178, 179, 180, 181, 183, 184, 185, 186],
        'mtq-plus.html': [191, 193, 195, 196, 197, 198, 200, 201, 202, 203, 205, 206, 207, 209, 210],
        'terapia-dzwiekiem.html': [215, 217, 219, 220, 221, 222, 224, 225, 226, 228, 229],
        'warsztaty-i-szkolenia.html': [233, 235, 239, 240, 241, 242, 243, 245, 246, 247, 249, 250, 252, 253, 254],
    }
    # PDF markers/commands are layout instructions, except the explicitly retained
    # biography placeholders. PDF list numbers become native HTML list markers.
    for page, indices in ranges.items():
        main = next(n for n in Document((ROOT / page).read_text()).nodes if n.tag == 'main')
        actual = tokens(main.text())
        for i in indices:
            expected = corrected(blocks[i]).replace('<COACHING>', '')
            expected = re.sub(r'<link:?\s+https://[^>]+>', '', expected)
            expected = re.sub(r'\s*–\s*wpisane w excela\s*$', '', expected)
            expected = re.sub(r'\b[1-9]\.\s+', '', expected).replace('·', ' ').replace('●', ' ')
            assert tokens(expected) in actual, (page, i, expected)
    print(f'OK — {sum(map(len, ranges.values()))} bloków briefu: wszystkie akapity, nagłówki, listy i placeholdery; tylko jawne korekty.')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source-dir', type=Path, default=ROOT / 'Wytyczne')
    args = parser.parse_args()
    audit(args.source_dir)
